"""
Gera Resumo_Cobertura_BTG.xlsx — planilha para compartilhar com o time.
Mostra quais categorias de CMV passam pelo BTG e quais fornecedores estão fora.
"""
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = r"C:\Users\Usuario\OneDrive\GiRafa_Code\Greenjoy\Resumo_Cobertura_BTG.xlsx"

# ── Dados BTG por categoria (jan-abr 2026) ───────────────────────────────────
btg_cat = [
    ("Hortaliças",              170_280.32,  9),
    ("Queijos",                 143_611.13,  4),
    ("Embalagens",              117_088.98, 12),
    ("Molhos Cozinha Central",  114_602.00,  1),  # reclassificado como Royalties
    ("Estoque Seco",             90_050.01, 10),
    ("Sobremesas e Smoothies",   72_140.10,  8),
    ("Proteínas",                55_884.89,  3),
    ("Tortas",                   22_161.60,  1),
    ("Bebidas",                  16_649.76,  4),  # Bebidas + bebidas merged
]

# ── Fornecedores fora do BTG ─────────────────────────────────────────────────
sem_btg = {
    "Bebidas":               ["Komfe", "PDV DISTRIBUICAO V A B LTDA", "SPON DISTRIBUIDORA DE BEBI"],
    "Embalagens":            ["BRUPACK EMBALAGENS E DESCARTAVEIS", "Buypack Brasil Descartaveis",
                              "Clay Transportes", "CONTAAZUL SOFTWARE LTDA",
                              "Embalagens Castropil", "GOOD PACK IND COM E REPR",
                              "JPL do Brasil", "Kraftbowls", "Lseki",
                              "Papeltec Embalagens", "Printbag", "RMRC EMBALAGENS EIRELI",
                              "Trenier"],
    "Estoque Seco":          ["Blend Alimentos", "COLAVITA BRASIL C I E EIRELI",
                              "EDUARDO JOSE DA SILVA DISTRIBUICAO", "F H O M IND C E P A LTDA",
                              "Freeway", "JTC DISTRIBUIDORA LTDA.", "MADÁ FOODS IND E COMERCIO",
                              "SEQUOIA ALIMENTOS LTDA EPP", "SUPERFOOD COMERCIO DE ALIMENTOS",
                              "SYNTHESIZE NUTRICAO HUMANA LTDA", "YAPAY PAGAMENTOS ONLINE", "Zaravia"],
    "Hortaliças":            ["ADRIANA HITOMI KOBAYASHI LIN", "ASOLUM AGRICULTURA TECNOLOGICA",
                              "Avocado Jaguacy", "CIUFFI HORTIFRUTI EIRELI", "GOOD FARMY",
                              "Henry Kim Young Lin", "MANIA DE AVOCADO LTDA",
                              "MIMOS DISTRIBUIDORA", "OFR COMERCIAL ANALISE DE CREDITO",
                              "RAFAEL DOS SANTOS SILVA", "THIAGO PEREIRA DAYNEZ"],
    "Proteínas":             ["APETITO FOODS LTDA", "BARON ALIMENTARE LTDA",
                              "Casa De Carnes Charm Da Be", "DISTRIBUIDORA DE CARNES SISSA",
                              "EMPORIO MEGA 100 COMERCIO DE", "FENIX FOODS ALIMENTOS EIRELI",
                              "Flamboia Alimentos", "FRIGORIFICO COWPIG LTDA.",
                              "FRUMAR FRUTOS DO MAR LTDA", "IDRA Pescados", "JBS S.A",
                              "M.P.F. NOVA UNIAO ALIMENTOS EIRELI", "MFISH COMERCIO DE PESCADOS",
                              "MINERVA FUNDO DE INVESTIMENTO", "MULTIFRANGOS COM ALIM LTDA ME",
                              "NAF COMERCIAL DE ALIMENTOS (Carnes)", "RDD COMERCIO DE ALIMENTOS",
                              "SANSUL ALIMENTOS LTDA", "TUNA IND E COM DE PESCADOS",
                              "3GX"],
    "Queijos":               ["AEB NEGOCIOS E CONEXOES LTDA", "ALLFOOD IMPORTACAO IND E COMERCIO",
                              "BUFALISSIMA", "COMERCIAL REJAN FOODS LTDA",
                              "DMART", "DULLER COM DE FRIOS E DERIVADOS",
                              "4F DISTRIBUIDORA DE ALIMENTOS EIRELI", "J V C D LATICINIOS LTDA",
                              "NAF COMERCIAL DE ALIMENTOS (Queijos)", "QUEIJOS FINOS", "Transleite"],
    "Sobremesas e Smoothies":["Agua Coco Buturi", "Ana Maria de Carvalho",
                              "EVOLAT LATICINIOS VEGETAIS LTD", "Frooty", "Marcos Monica",
                              "PAGHIPER SERVICOS ONLINE", "PUSH MATCHA COMERCIO DE PRODUTOS",
                              "SUBLYME FABRICACAO E COMERCIO DE ALIMENTOS"],
    "Molhos Cozinha Central":[], "Tortas":[],
}

# ── Montar df resumo ─────────────────────────────────────────────────────────
rows = []
total_btg = sum(v for _, v, _ in btg_cat if _ != 1 or True)  # todos
for cat, valor, forn_ativos in btg_cat:
    fora = sem_btg.get(cat, [])
    obs = "Reclassificado como Royalties (Despesa)" if cat == "Molhos Cozinha Central" else ""
    rows.append({
        "Categoria":             cat,
        "Valor BTG Jan-Abr (R$)": valor,
        "Fornecedores com BTG":  forn_ativos,
        "Fornecedores sem BTG":  len(fora),
        "Observação":            obs,
    })

df_res = pd.DataFrame(rows)
df_res["Total fornecedores"] = df_res["Fornecedores com BTG"] + df_res["Fornecedores sem BTG"]
df_res["% cobertura BTG"] = (df_res["Fornecedores com BTG"] / df_res["Total fornecedores"]).fillna(1)

# Reordenar colunas
df_res = df_res[["Categoria", "Valor BTG Jan-Abr (R$)",
                  "Fornecedores com BTG", "Fornecedores sem BTG",
                  "Total fornecedores", "% cobertura BTG", "Observação"]]

# Totais
total_row = {
    "Categoria": "TOTAL",
    "Valor BTG Jan-Abr (R$)": df_res["Valor BTG Jan-Abr (R$)"].sum(),
    "Fornecedores com BTG":   df_res["Fornecedores com BTG"].sum(),
    "Fornecedores sem BTG":   df_res["Fornecedores sem BTG"].sum(),
    "Total fornecedores":     df_res["Total fornecedores"].sum(),
    "% cobertura BTG":        df_res["Fornecedores com BTG"].sum() / df_res["Total fornecedores"].sum(),
    "Observação": "",
}
df_res = pd.concat([df_res, pd.DataFrame([total_row])], ignore_index=True)

# ── Montar df detalhe sem BTG ────────────────────────────────────────────────
det_rows = []
for cat, forn_list in sem_btg.items():
    for f in forn_list:
        det_rows.append({"Categoria": cat, "Fornecedor": f,
                         "Canal provável": "Banco Inter / outro"
                         if cat == "Proteínas" else "Verificar"})
df_det = pd.DataFrame(det_rows).sort_values(["Categoria", "Fornecedor"])

# ── Escrever Excel ───────────────────────────────────────────────────────────
with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
    df_res.to_excel(xw, sheet_name="Resumo por Categoria", index=False)
    df_det.to_excel(xw, sheet_name="Fornecedores sem BTG", index=False)

    # ── Formatar aba Resumo ──────────────────────────────────────────────────
    ws = xw.sheets["Resumo por Categoria"]

    VERDE  = PatternFill("solid", fgColor="C6EFCE")
    AMARELO= PatternFill("solid", fgColor="FFEB9C")
    LARANJA= PatternFill("solid", fgColor="FFCC99")
    HEADER = PatternFill("solid", fgColor="1F3864")
    TOTAL  = PatternFill("solid", fgColor="D9D9D9")

    header_font = Font(bold=True, color="FFFFFF")
    total_font  = Font(bold=True)
    thin = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin")
    )

    # Header
    for cell in ws[1]:
        cell.fill = HEADER
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    n_rows = len(df_res)
    pct_col = 6  # coluna F = % cobertura BTG

    for row_idx in range(2, n_rows + 2):
        is_total = row_idx == n_rows + 1
        pct_cell = ws.cell(row=row_idx, column=pct_col)
        pct_cell.number_format = "0%"

        for col_idx in range(1, len(df_res.columns) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin
            if is_total:
                cell.fill = TOTAL
                cell.font = total_font
            elif col_idx == 2:
                cell.number_format = '#,##0.00'

        if not is_total:
            pct = ws.cell(row=row_idx, column=pct_col).value or 0
            if isinstance(pct, float):
                fill = VERDE if pct >= 0.8 else (AMARELO if pct >= 0.5 else LARANJA)
                for col_idx in range(1, len(df_res.columns) + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = fill

    # Largura das colunas
    widths = [28, 22, 22, 22, 20, 18, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ── Formatar aba Detalhe ─────────────────────────────────────────────────
    ws2 = xw.sheets["Fornecedores sem BTG"]
    for cell in ws2[1]:
        cell.fill = HEADER
        cell.font = header_font
    for col, w in zip(["A", "B", "C"], [28, 45, 25]):
        ws2.column_dimensions[col].width = w
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.border = thin

print(f"Gerado: {OUT}")
