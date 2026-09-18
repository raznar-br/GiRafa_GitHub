"""
adicionar_aba_embalagens.py — Adiciona aba de detalhe de Embalagens ao Excel de auditoria.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date

ARQUIVO = r"C:\Users\Usuario\OneDrive\GiRafa_Code\Greenjoy\Auditoria_KPIs_Dashboard.xlsx"

# Paleta
VERDE_ESCURO   = "1a6b3c"
VERDE_CLARO    = "dcfce7"
AMARELO        = "fef9c3"
AMARELO_ESCURO = "ca8a04"
VERMELHO       = "fee2e2"
VERMELHO_ESCURO= "dc2626"
CINZA          = "f1f5f9"
CINZA_BORDA    = "cbd5e1"
BRANCO         = "ffffff"
PRETO          = "0f172a"
AZUL_CLARO     = "eff6ff"
AZUL_ESCURO    = "1d4ed8"

BRL = '#,##0.00'
PCT = '0.0%'

def thin_border():
    s = Side(style="thin", color=CINZA_BORDA)
    return Border(left=s, right=s, top=s, bottom=s)

def hdr(ws, row, col, value, fg=BRANCO, bg=VERDE_ESCURO, bold=True, sz=10):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name="Calibri", bold=bold, size=sz, color=fg)
    c.fill = PatternFill("solid", fgColor=bg)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = thin_border()
    return c

def cell(ws, row, col, value, bold=False, color=PRETO, bg=BRANCO, align="left", fmt=None, sz=10):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name="Calibri", bold=bold, size=sz, color=color)
    c.fill = PatternFill("solid", fgColor=bg)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=False)
    c.border = thin_border()
    if fmt:
        c.number_format = fmt
    return c

def titulo(ws, texto, sub=""):
    ws.merge_cells("A1:I1")
    c = ws["A1"]
    c.value = texto
    c.font = Font(name="Calibri", bold=True, size=14, color=BRANCO)
    c.fill = PatternFill("solid", fgColor=VERDE_ESCURO)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30
    if sub:
        ws.merge_cells("A2:I2")
        c2 = ws["A2"]
        c2.value = sub
        c2.font = Font(name="Calibri", italic=True, size=10, color=PRETO)
        c2.fill = PatternFill("solid", fgColor=CINZA)
        c2.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[2].height = 16

# ── Dados: resumo por fornecedor/mês (dez/25 → mai/26)
resumo = [
    # (mes, fornecedor, fornecedor_grupo, qtd, total)
    ("2026-05","TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA","Tuicial Grafica",1,3102.85),
    ("2026-05","FNS COMERCIO E IMPORTACAO LTDA","FNS Comercio",1,2839.00),
    ("2026-05","UNIPEL IMPORTACAO E I E LTDA","UNIPEL",2,2211.64),
    ("2026-05","PAPEL. PLASTICO ITUPEVA LTDA","Papel Plástico Itupeva",1,990.31),
    ("2026-05","GS COMERCIO DE PAPEIS","GS Comércio de Papéis",1,910.00),
    ("2026-05","PREMIUM PRODUTOS DESCARTAVEIS LTDA","Premium Descartáveis",1,369.06),
    ("2026-05","LEO ETIQUETAS E ROTULOS ADESIVOS LT","Leo Etiquetas",1,336.30),
    ("2026-05","MEIWA INDUSTRIA COMERCIO LTDA","Meiwa",1,247.72),
    ("2026-05","GJ FUNCHAL LTDA","GJ Funchal",1,150.00),
    ("2026-04","FNS COMERCIO E IMPORTACAO LTDA","FNS Comercio",4,11210.04),
    ("2026-04","UNIPEL IMPORTACAO E I E LTDA","UNIPEL",5,4435.66),
    ("2026-04","TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA","Tuicial Grafica",1,3102.85),
    ("2026-04","DM7 BRASIL INDUSTRIA E COMERCIO DE EMBALAGENS LTDA","DM7 Brasil",1,1126.00),
    ("2026-04","PAPEL. PLASTICO ITUPEVA LTDA","Papel Plástico Itupeva",2,1061.02),
    ("2026-04","GS COMERCIO DE PAPEIS","GS Comércio de Papéis",1,910.00),
    ("2026-04","VM PACK SOLUCOES EM EMBALAGENS","VM Pack",1,665.00),
    ("2026-04","GJ FUNCHAL LTDA","GJ Funchal",2,337.52),
    ("2026-03","FNS COMERCIO E IMPORTACAO LTDA","FNS Comercio",3,6167.46),
    ("2026-03","TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA","Tuicial Grafica",2,4650.32),
    ("2026-03","UNIPEL IMPORTACAO E I E LTDA","UNIPEL",3,2353.08),
    ("2026-03","PAPEL. PLASTICO ITUPEVA LTDA","Papel Plástico Itupeva",3,1591.73),
    ("2026-03","GS COMERCIO DE PAPEIS","GS Comércio de Papéis",1,875.00),
    ("2026-03","MARIA DIVINA AEROPIPA FIOS L L","Maria Divina (fios/amarrilhos)",1,542.58),
    ("2026-03","PREMIUM PAPEIS PERSONALIZADOS LTDA","Premium Descartáveis",1,528.65),
    ("2026-03","LEO ETIQUETAS E ROTULOS ADESIVOS LT","Leo Etiquetas",1,336.30),
    ("2026-03","MEIWA INDUSTRIA COMERCIO LTDA","Meiwa",1,297.80),
    ("2026-03","GJ FUNCHAL LTDA","GJ Funchal",1,37.63),
    ("2026-02","FNS COMERCIO E IMPORTACAO LTDA","FNS Comercio",4,7857.98),
    ("2026-02","UNIPEL IMPORTACAO E I E LTDA","UNIPEL",4,2617.41),
    ("2026-02","TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA","Tuicial Grafica",1,2325.16),
    ("2026-02","PAPEL. PLASTICO ITUPEVA LTDA","Papel Plástico Itupeva",3,1997.65),
    ("2026-02","LEO ETIQUETAS E ROTULOS ADESIVOS LT","Leo Etiquetas",1,594.00),
    ("2026-02","MARIA DIVINA AEROPIPA FIOS L L","Maria Divina (fios/amarrilhos)",1,542.58),
    ("2026-02","MEIWA INDUSTRIA COMERCIO LTDA","Meiwa",1,218.12),
    ("2026-02","GJ FUNCHAL LTDA","GJ Funchal",1,37.63),
    ("2026-01","FNS COMERCIO E IMPORTACAO LTDA","FNS Comercio",3,4997.98),
    ("2026-01","UNIPEL IMPORTACAO E I E LTDA","UNIPEL",3,3319.01),
    ("2026-01","TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA","Tuicial Grafica",1,2325.16),
    ("2026-01","PREMIUM PRODUTOS DESCARTAVEIS LTDA","Premium Descartáveis",2,1328.49),
    ("2026-01","DM7 BRASIL INDUSTRIA E COMERCIO DE EMBALAGENS LTDA","DM7 Brasil",1,1126.00),
    ("2026-01","GS COMERCIO DE PAPEIS","GS Comércio de Papéis",1,1050.00),
    ("2026-01","LEO ETIQUETAS E ROTULOS ADESIVOS LT","Leo Etiquetas",1,704.00),
    ("2026-01","VM PACK SOLUCOES EM EMBALAGENS LTD","VM Pack",1,660.00),
    ("2026-01","MEIWA INDUSTRIA COMERCIO LTDA","Meiwa",1,232.59),
    ("2026-01","GJ FUNCHAL LTDA","GJ Funchal",1,37.63),
    ("2025-12","FNS COMERCIO E IMPORTACAO LTDA","FNS Comercio",4,8856.02),
    ("2025-12","TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA","Tuicial Grafica",2,4591.72),
    ("2025-12","UNIPEL IMPORTACAO E I E LTDA","UNIPEL",4,3437.77),
    ("2025-12","PAPEL. PLASTICO ITUPEVA LTDA","Papel Plástico Itupeva",4,1513.16),
    ("2025-12","LEO ETIQUETAS E ROTULOS ADESIVOS LT","Leo Etiquetas",2,881.00),
    ("2025-12","MARIA DIVINA AEROPIPA FIOS L L","Maria Divina (fios/amarrilhos)",1,542.58),
    ("2025-12","MEIWA INDUSTRIA COMERCIO LTDA","Meiwa",1,232.59),
    ("2025-12","GJ FUNCHAL LTDA","GJ Funchal",1,37.63),
]

# ── Transações individuais (dez/25 → mai/26)
transacoes = [
    ("2026-05-18","2026-05","GS COMERCIO DE PAPEIS","Pagamento de boleto enviado para GS COMERCIO DE PAPEIS",910.00),
    ("2026-05-13","2026-05","FNS COMERCIO E IMPORTACAO LTDA","Pagamento de boleto enviado para FNS COMERCIO E IMPORTACAO LTDA",2839.00),
    ("2026-05-13","2026-05","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",1379.43),
    ("2026-05-11","2026-05","PREMIUM PRODUTOS DESCARTAVEIS LTDA","Pagamento de boleto enviado para PREMIUM PRODUTOS DESCARTAVEIS LTDA",369.06),
    ("2026-05-11","2026-05","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",832.21),
    ("2026-05-11","2026-05","GJ FUNCHAL LTDA","Pix enviado para Gj Funchal Ltda",150.00),
    ("2026-05-08","2026-05","MEIWA INDUSTRIA COMERCIO LTDA","Pagamento de boleto enviado para MEIWA INDUSTRIA COMERCIO LTDA",247.72),
    ("2026-05-07","2026-05","TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA","Pagamento de boleto enviado para TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA",3102.85),
    ("2026-05-06","2026-05","LEO ETIQUETAS E ROTULOS ADESIVOS LT","Pagamento de boleto enviado para LEO ETIQUETAS E ROTULOS ADESIVOS LT",336.30),
    ("2026-05-05","2026-05","PAPEL. PLASTICO ITUPEVA LTDA","Pagamento de boleto enviado para PAPEL PLASTICO ITUPEVA LTDA",990.31),
    ("2026-04-29","2026-04","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",1113.04),
    ("2026-04-29","2026-04","GJ FUNCHAL LTDA","Pix enviado para Gj Funchal Ltda",37.82),
    ("2026-04-29","2026-04","FNS COMERCIO E IMPORTACAO LTDA","Pagamento de boleto enviado para FNS COMERCIO E IMPORTACAO LTDA",3079.01),
    ("2026-04-28","2026-04","GJ FUNCHAL LTDA","Pix enviado para Gj Funchal Ltda",299.70),
    ("2026-04-22","2026-04","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",454.31),
    ("2026-04-20","2026-04","VM PACK SOLUCOES EM EMBALAGENS","Pagamento de boleto enviado para VM PACK SOLUCOES EM EMBALAGENS",665.00),
    ("2026-04-16","2026-04","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",693.85),
    ("2026-04-15","2026-04","FNS COMERCIO E IMPORTACAO LTDA","Pagamento de boleto enviado para FNS COMERCIO E IMPORTACAO LTDA",3703.02),
    ("2026-04-14","2026-04","PAPEL. PLASTICO ITUPEVA LTDA","Pagamento de boleto enviado para PAPEL PLASTICO ITUPEVA LTDA",530.51),
    ("2026-04-10","2026-04","TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA","Pagamento de boleto enviado para TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA",3102.85),
    ("2026-04-08","2026-04","FNS COMERCIO E IMPORTACAO LTDA","Pagamento de boleto enviado para FNS COMERCIO E IMPORTACAO LTDA",1444.01),
    ("2026-04-08","2026-04","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",467.73),
    ("2026-04-07","2026-04","DM7 BRASIL INDUSTRIA E COMERCIO DE EMBALAGENS LTDA","Pagamento de boleto enviado para DM7 BRASIL INDUSTRIA E COMERCIO DE EMBALAGENS LTDA",1126.00),
    ("2026-04-06","2026-04","PAPEL. PLASTICO ITUPEVA LTDA","Pagamento de boleto enviado para PAPEL PLASTICO ITUPEVA LTDA",530.51),
    ("2026-04-02","2026-04","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",1706.73),
    ("2026-04-01","2026-04","FNS COMERCIO E IMPORTACAO LTDA","Pagamento de boleto enviado para FNS COMERCIO E IMPORTACAO LTDA",2984.00),
    ("2026-04-01","2026-04","GS COMERCIO DE PAPEIS","Pagamento de boleto enviado para GS COMERCIO DE PAPEIS",910.00),
    ("2026-03-31","2026-03","LEO ETIQUETAS E ROTULOS ADESIVOS LT","Pagamento de boleto enviado para LEO ETIQUETAS E ROTULOS ADESIVOS LT",336.30),
    ("2026-03-30","2026-03","PAPEL. PLASTICO ITUPEVA LTDA","Pagamento de boleto enviado para PAPEL PLASTICO ITUPEVA LTDA",733.67),
    ("2026-03-27","2026-03","GJ FUNCHAL LTDA","Pix enviado para Gj Funchal Ltda",37.63),
    ("2026-03-25","2026-03","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",492.51),
    ("2026-03-25","2026-03","MARIA DIVINA AEROPIPA FIOS L L","Pagamento de boleto enviado para MARIA DIVINA AEROPIPA FIOS L L",542.58),
    ("2026-03-24","2026-03","PREMIUM PAPEIS PERSONALIZADOS LTDA","Pagamento de boleto enviado para PREMIUM PAPEIS PERSONALIZADOS LTDA",528.65),
    ("2026-03-19","2026-03","TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA","Pagamento de boleto enviado para TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA",2325.16),
    ("2026-03-19","2026-03","TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA","Pagamento de boleto enviado para TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA",2325.16),
    ("2026-03-18","2026-03","PAPEL. PLASTICO ITUPEVA LTDA","Pagamento de boleto enviado para PAPEL PLASTICO ITUPEVA LTDA",530.51),
    ("2026-03-18","2026-03","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",903.44),
    ("2026-03-18","2026-03","FNS COMERCIO E IMPORTACAO LTDA","Pagamento de boleto enviado para FNS COMERCIO E IMPORTACAO LTDA",1444.01),
    ("2026-03-11","2026-03","MEIWA INDUSTRIA COMERCIO LTDA","Pagamento de boleto enviado para MEIWA INDUSTRIA COMERCIO LTDA",297.80),
    ("2026-03-11","2026-03","FNS COMERCIO E IMPORTACAO LTDA","Pagamento de boleto enviado para FNS COMERCIO E IMPORTACAO LTDA",2207.44),
    ("2026-03-04","2026-03","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",957.13),
    ("2026-03-04","2026-03","FNS COMERCIO E IMPORTACAO LTDA","Pagamento de boleto enviado para FNS COMERCIO E IMPORTACAO LTDA",2516.01),
    ("2026-03-04","2026-03","GS COMERCIO DE PAPEIS","Pagamento de boleto enviado para GS COMERCIO DE PAPEIS",875.00),
    ("2026-03-02","2026-03","PAPEL. PLASTICO ITUPEVA LTDA","Pagamento de boleto enviado para PAPEL PLASTICO ITUPEVA LTDA",327.55),
    ("2026-02-27","2026-02","LEO ETIQUETAS E ROTULOS ADESIVOS LT","Pagamento de boleto enviado para LEO ETIQUETAS E ROTULOS ADESIVOS LT",594.00),
    ("2026-02-27","2026-02","GJ FUNCHAL LTDA","Pix enviado para Gj Funchal Ltda",37.63),
    ("2026-02-25","2026-02","FNS COMERCIO E IMPORTACAO LTDA","Pagamento de boleto enviado para FNS COMERCIO E IMPORTACAO LTDA",1464.99),
    ("2026-02-25","2026-02","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",480.12),
    ("2026-02-18","2026-02","PAPEL. PLASTICO ITUPEVA LTDA","Pagamento de boleto enviado para PAPEL PLASTICO ITUPEVA LTDA",530.51),
    ("2026-02-18","2026-02","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",669.07),
    ("2026-02-18","2026-02","FNS COMERCIO E IMPORTACAO LTDA","Pagamento de boleto enviado para FNS COMERCIO E IMPORTACAO LTDA",2152.00),
    ("2026-02-12","2026-02","MEIWA INDUSTRIA COMERCIO LTDA","Pagamento de boleto enviado para MEIWA INDUSTRIA COMERCIO LTDA",218.12),
    ("2026-02-11","2026-02","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",475.99),
    ("2026-02-11","2026-02","FNS COMERCIO E IMPORTACAO LTDA","Pagamento de boleto enviado para FNS COMERCIO E IMPORTACAO LTDA",1464.99),
    ("2026-02-09","2026-02","PAPEL. PLASTICO ITUPEVA LTDA","Pagamento de boleto enviado para PAPEL PLASTICO ITUPEVA LTDA",936.63),
    ("2026-02-06","2026-02","MARIA DIVINA AEROPIPA FIOS L L","Pagamento de boleto enviado para MARIA DIVINA AEROPIPA FIOS L L",542.58),
    ("2026-02-06","2026-02","TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA","Pagamento de boleto enviado para TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA",2325.16),
    ("2026-02-05","2026-02","FNS COMERCIO E IMPORTACAO LTDA","Pagamento de boleto enviado para FNS COMERCIO E IMPORTACAO LTDA",2776.00),
    ("2026-02-04","2026-02","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",992.23),
    ("2026-02-03","2026-02","PAPEL. PLASTICO ITUPEVA LTDA","Pagamento de boleto enviado para PAPEL PLASTICO ITUPEVA LTDA",530.51),
    ("2026-01-30","2026-01","GJ FUNCHAL LTDA","Pix enviado para Gj Funchal Ltda",37.63),
    ("2026-01-23","2026-01","DM7 BRASIL INDUSTRIA E COMERCIO DE EMBALAGENS LTDA","Pagamento de boleto enviado para DM7 BRASIL INDUSTRIA E COMERCIO DE EMBALAGENS LTDA",1126.00),
    ("2026-01-20","2026-01","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",1345.35),
    ("2026-01-19","2026-01","PREMIUM PRODUTOS DESCARTAVEIS LTDA","Pagamento de boleto enviado para PREMIUM PRODUTOS DESCARTAVEIS LTDA",1002.37),
    ("2026-01-14","2026-01","FNS COMERCIO E IMPORTACAO LTDA","Pagamento de boleto enviado para FNS COMERCIO E IMPORTACAO LTDA",2131.01),
    ("2026-01-09","2026-01","TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA","Pagamento de boleto enviado para TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA",2325.16),
    ("2026-01-07","2026-01","FNS COMERCIO E IMPORTACAO LTDA","Pagamento de boleto enviado para FNS COMERCIO E IMPORTACAO LTDA",1401.98),
    ("2026-01-07","2026-01","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",1263.78),
    ("2026-01-07","2026-01","LEO ETIQUETAS E ROTULOS ADESIVOS LT","Pagamento de boleto enviado para LEO ETIQUETAS E ROTULOS ADESIVOS LT",704.00),
    ("2026-01-06","2026-01","PREMIUM PRODUTOS DESCARTAVEIS LTDA","Pagamento de boleto enviado para PREMIUM PRODUTOS DESCARTAVEIS LTDA",326.12),
    ("2026-01-05","2026-01","VM PACK SOLUCOES EM EMBALAGENS LTD","Pagamento de boleto enviado para VM PACK SOLUCOES EM EMBALAGENS LTD",660.00),
    ("2026-01-05","2026-01","MEIWA INDUSTRIA COMERCIO LTDA","Pagamento de boleto enviado para MEIWA INDUSTRIA COMERCIO LTDA",232.59),
    ("2026-01-02","2026-01","GS COMERCIO DE PAPEIS","Pagamento de boleto enviado para GS COMERCIO DE PAPEIS",1050.00),
    ("2026-01-02","2026-01","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",709.88),
    ("2026-01-02","2026-01","FNS COMERCIO E IMPORTACAO LTDA","Pagamento de boleto enviado para FNS COMERCIO E IMPORTACAO LTDA",1464.99),
    ("2025-12-30","2025-12","GJ FUNCHAL LTDA","Pix enviado para Gj Funchal Ltda",37.63),
    ("2025-12-29","2025-12","PAPEL. PLASTICO ITUPEVA LTDA","Pagamento de boleto enviado para PAPEL PLASTICO ITUPEVA LTDA",327.55),
    ("2025-12-26","2025-12","LEO ETIQUETAS E ROTULOS ADESIVOS LT","Pagamento de boleto enviado para LEO ETIQUETAS E ROTULOS ADESIVOS LT",704.00),
    ("2025-12-24","2025-12","FNS COMERCIO E IMPORTACAO LTDA","Pagamento de boleto enviado para FNS COMERCIO E IMPORTACAO LTDA",1444.01),
    ("2025-12-24","2025-12","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",567.88),
    ("2025-12-16","2025-12","FNS COMERCIO E IMPORTACAO LTDA","Pagamento de boleto enviado para FNS COMERCIO E IMPORTACAO LTDA",2776.00),
    ("2025-12-16","2025-12","TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA","Pagamento de boleto enviado para TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA",2325.16),
    ("2025-12-16","2025-12","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",877.63),
    ("2025-12-15","2025-12","PAPEL. PLASTICO ITUPEVA LTDA","Pagamento de boleto enviado para PAPEL PLASTICO ITUPEVA LTDA",530.51),
    ("2025-12-12","2025-12","MARIA DIVINA AEROPIPA FIOS L L","Pagamento de boleto enviado para MARIA DIVINA AEROPIPA FIOS L L",542.58),
    ("2025-12-10","2025-12","FNS COMERCIO E IMPORTACAO LTDA","Pagamento de boleto enviado para FNS COMERCIO E IMPORTACAO LTDA",1444.01),
    ("2025-12-10","2025-12","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",1035.63),
    ("2025-12-09","2025-12","LEO ETIQUETAS E ROTULOS ADESIVOS LT","Pagamento de boleto enviado para LEO ETIQUETAS E ROTULOS ADESIVOS LT",177.00),
    ("2025-12-08","2025-12","PAPEL. PLASTICO ITUPEVA LTDA","Pagamento de boleto enviado para PAPEL PLASTICO ITUPEVA LTDA",327.55),
    ("2025-12-05","2025-12","TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA","Pagamento de boleto enviado para TUICIAL INDUSTRIA GRAFICA E EDITORA LTDA",2266.56),
    ("2025-12-04","2025-12","MEIWA INDUSTRIA COMERCIO LTDA","Pagamento de boleto enviado para MEIWA INDUSTRIA COMERCIO LTDA",232.59),
    ("2025-12-03","2025-12","FNS COMERCIO E IMPORTACAO LTDA","Pagamento de boleto enviado para FNS COMERCIO E IMPORTACAO LTDA",3192.00),
    ("2025-12-03","2025-12","UNIPEL IMPORTACAO E I E LTDA","Pagamento de boleto enviado para UNIPEL IMPORTACAO E I E LTDA",956.63),
    ("2025-12-01","2025-12","PAPEL. PLASTICO ITUPEVA LTDA","Pagamento de boleto enviado para PAPEL PLASTICO ITUPEVA LTDA",327.55),
]

# Fornecedores que merecem atenção / validação
ALERTAS = {
    "MARIA DIVINA AEROPIPA FIOS L L": "Verificar — fios e linhas (amarrilhos para embalagem?)",
    "GJ FUNCHAL LTDA": "Verificar — valores muito pequenos (R$37-300), o que é?",
    "CONTAAZUL SOFTWARE LTDA": "ALERTA — software contábil, não é embalagem",
    "CLAY TRANSPORTES": "ALERTA — transportadora, não é embalagem",
    "MEGA BANK SECURITIZADORA DE ATIVOS EMPRE": "ALERTA — financeira (boleto de embalagem via banco?)",
}

def main():
    wb = openpyxl.load_workbook(ARQUIVO)

    # Remover aba anterior se já existir
    if "7. Embalagens — Detalhe" in wb.sheetnames:
        del wb["7. Embalagens — Detalhe"]

    ws = wb.create_sheet("7. Embalagens — Detalhe")

    titulo(ws, "Embalagens — Detalhe de Gastos (BTG, dez/25–mai/26)",
           "Fonte: btg_movimentacoes × jafeb_despesas_mapping WHERE categoria_2='Embalagens'. "
           "Abr/26 = R$22.848 (10,5% do CMV). Validar cada fornecedor.")

    # ── SEÇÃO 1: Resumo por fornecedor (totais agregados no período)
    ws.merge_cells("A4:H4")
    sh = ws["A4"]
    sh.value = "TOTAIS POR FORNECEDOR — dez/25 a mai/26"
    sh.font = Font(name="Calibri", bold=True, size=11, color=BRANCO)
    sh.fill = PatternFill("solid", fgColor=AZUL_ESCURO)
    sh.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[4].height = 20

    # Agregar por fornecedor
    por_forn = {}
    for mes, forn, grupo, qtd, total in resumo:
        if forn not in por_forn:
            por_forn[forn] = {"grupo": grupo, "qtd": 0, "total": 0.0, "meses": set()}
        por_forn[forn]["qtd"] += qtd
        por_forn[forn]["total"] += total
        por_forn[forn]["meses"].add(mes)

    total_geral = sum(v["total"] for v in por_forn.values())
    forn_sorted = sorted(por_forn.items(), key=lambda x: -x[1]["total"])

    h_forn = ["Fornecedor (BTG)", "Grupo/Apelido", "Pagamentos", "Total R$", "% do total", "Meses ativos", "Alerta / Validar"]
    w_forn = [45, 30, 12, 15, 12, 12, 45]
    for i, (h, w) in enumerate(zip(h_forn, w_forn), 1):
        hdr(ws, 5, i, h, fg=BRANCO, bg=AZUL_ESCURO)
        ws.column_dimensions[get_column_letter(i)].width = w

    for row_i, (forn, v) in enumerate(forn_sorted, 6):
        alerta = ALERTAS.get(forn, "")
        bg = AMARELO if alerta.startswith("Verificar") else VERMELHO if alerta.startswith("ALERTA") else BRANCO
        pct_v = v["total"] / total_geral if total_geral else 0
        cell(ws, row_i, 1, forn,              bold=True, bg=bg)
        cell(ws, row_i, 2, v["grupo"],         bg=bg)
        cell(ws, row_i, 3, v["qtd"],           align="center", bg=bg)
        cell(ws, row_i, 4, v["total"],         align="right", fmt=BRL, bg=bg, bold=True)
        cell(ws, row_i, 5, pct_v,              align="right", fmt=PCT, bg=bg)
        cell(ws, row_i, 6, len(v["meses"]),    align="center", bg=bg)
        cell(ws, row_i, 7, alerta,             bg=bg,
             color=VERMELHO_ESCURO if alerta.startswith("ALERTA") else
                   AMARELO_ESCURO if alerta.startswith("Verificar") else PRETO,
             bold=bool(alerta))
        ws.row_dimensions[row_i].height = 16

    # Total
    r_tot = len(forn_sorted) + 6
    cell(ws, r_tot, 1, "TOTAL", bold=True, bg=CINZA)
    cell(ws, r_tot, 2, "", bg=CINZA)
    cell(ws, r_tot, 3, sum(v["qtd"] for v in por_forn.values()), align="center", bold=True, bg=CINZA)
    cell(ws, r_tot, 4, total_geral, align="right", fmt=BRL, bold=True, bg=CINZA)
    cell(ws, r_tot, 5, 1.0, align="right", fmt=PCT, bg=CINZA)
    ws.row_dimensions[r_tot].height = 18

    # ── SEÇÃO 2: Resumo mensal
    r_sec2 = r_tot + 2
    ws.merge_cells(f"A{r_sec2}:H{r_sec2}")
    s2 = ws[f"A{r_sec2}"]
    s2.value = "TOTAL POR MÊS"
    s2.font = Font(name="Calibri", bold=True, size=11, color=BRANCO)
    s2.fill = PatternFill("solid", fgColor=AZUL_ESCURO)
    s2.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[r_sec2].height = 20

    h_mes = ["Mês", "Total Embalagens", "% sobre CMV mensal*"]
    cmv_por_mes = {"2025-12": 246545.94, "2026-01": 187583.92, "2026-02": 206649.60,
                   "2026-03": 200349.14, "2026-04": 216703.58, "2026-05": 110469.30}
    for i, h in enumerate(h_mes, 1):
        hdr(ws, r_sec2 + 1, i, h, fg=BRANCO, bg=AZUL_ESCURO)

    por_mes = {}
    for mes, forn, grupo, qtd, total in resumo:
        por_mes[mes] = por_mes.get(mes, 0.0) + total

    for row_i, mes in enumerate(sorted(por_mes.keys(), reverse=True), r_sec2 + 2):
        tot_m = por_mes[mes]
        cmv_m = cmv_por_mes.get(mes, 0)
        pct_cmv = tot_m / cmv_m if cmv_m else 0
        cell(ws, row_i, 1, mes,     bold=True, align="center")
        cell(ws, row_i, 2, tot_m,   align="right", fmt=BRL, bold=True)
        cell(ws, row_i, 3, pct_cmv, align="right", fmt=PCT,
             color=VERMELHO_ESCURO if pct_cmv > 0.12 else AMARELO_ESCURO if pct_cmv > 0.09 else PRETO)
        ws.row_dimensions[row_i].height = 16

    nota_r = r_sec2 + 2 + len(por_mes)
    ws.merge_cells(f"A{nota_r}:H{nota_r}")
    cn = ws[f"A{nota_r}"]
    cn.value = "* % calculado sobre total CMV do mês (BTG). Referência: Abr/26 = R$22.848 / R$216.703 = 10,5%"
    cn.font = Font(name="Calibri", italic=True, size=9, color=PRETO)
    cn.fill = PatternFill("solid", fgColor=CINZA)
    cn.alignment = Alignment(horizontal="left")
    ws.row_dimensions[nota_r].height = 14

    # ── SEÇÃO 3: Transações individuais
    r_sec3 = nota_r + 2
    ws.merge_cells(f"A{r_sec3}:H{r_sec3}")
    s3 = ws[f"A{r_sec3}"]
    s3.value = "TRANSAÇÕES INDIVIDUAIS — dez/25 a mai/26 (do mais recente ao mais antigo)"
    s3.font = Font(name="Calibri", bold=True, size=11, color=BRANCO)
    s3.fill = PatternFill("solid", fgColor=AZUL_ESCURO)
    s3.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[r_sec3].height = 20

    h_tx = ["Data", "Mês", "Fornecedor", "Descrição BTG", "Valor R$", "Alerta"]
    w_tx = [12, 10, 45, 55, 15, 40]
    for i, (h, w) in enumerate(zip(h_tx, w_tx), 1):
        hdr(ws, r_sec3 + 1, i, h, fg=BRANCO, bg=AZUL_ESCURO)
        ws.column_dimensions[get_column_letter(i)].width = max(
            ws.column_dimensions[get_column_letter(i)].width, w
        )

    cores_mes = {}
    palette = [BRANCO, AZUL_CLARO]
    mes_idx = 0
    last_mes = None

    for row_i, (data, mes, forn, descr, valor) in enumerate(transacoes, r_sec3 + 2):
        if mes != last_mes:
            last_mes = mes
            mes_idx = (mes_idx + 1) % 2
        bg = palette[mes_idx]
        alerta = ALERTAS.get(forn, "")
        if alerta:
            bg = AMARELO if alerta.startswith("Verificar") else VERMELHO
        cell(ws, row_i, 1, data,   align="center", bg=bg)
        cell(ws, row_i, 2, mes,    align="center", bg=bg)
        cell(ws, row_i, 3, forn,   bg=bg)
        cell(ws, row_i, 4, descr,  bg=bg, sz=9)
        cell(ws, row_i, 5, valor,  align="right", fmt=BRL, bold=True, bg=bg)
        cell(ws, row_i, 6, alerta, bg=bg,
             color=VERMELHO_ESCURO if alerta.startswith("ALERTA") else
                   AMARELO_ESCURO if alerta.startswith("Verificar") else PRETO,
             sz=9)
        ws.row_dimensions[row_i].height = 15

    ws.freeze_panes = f"A{r_sec3 + 2}"

    wb.save(ARQUIVO)
    print(f"Aba adicionada. Arquivo salvo: {ARQUIVO}")


if __name__ == "__main__":
    main()
